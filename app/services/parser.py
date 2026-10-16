
"""
文件名：app/services/parser.py
作用：多格式文档解析（流水线的"拆包员"）。
      把 PDF / DOCX / MD / TXT / XLSX 统一解析成 LangChain Document 列表。
对应规范：SPEC.md §8.1 文档解析
"""

# ---- 导入工具：每一行对应一个"格式读取器" ----
from pathlib import Path          # 把路径字符串变成路径对象，跨平台更稳妥

import pymupdf                    # PyMuPDF：读 PDF 的专业库
from docx import Document as WordDocument   # python-docx：读 Word，改名 WordDocument 避免冲突
from openpyxl import load_workbook          # openpyxl：读 Excel
from langchain_core.documents import Document  # LangChain 的"统一卡片"类


def parse_file(file_path: str | Path) -> list[Document]:
    """
    输入：一个文件的路径（字符串或 Path 对象都行）
    输出：Document 列表。每张卡片 = 正文(page_content) + 出处标签(metadata)
    """
    path = Path(file_path)
    suffix = path.suffix.lower()      # 取文件后缀并转小写，如 ".PDF" → ".pdf"
    documents = []                    # 准备一个空列表，装好的卡片往里放

    # ============ 情况 1：PDF 文件，逐页读 ============
    if suffix == ".pdf":
        # with 语句：用完自动关闭文件，不用手动记得关
        with pymupdf.open(path) as pdf:
            # enumerate：同时拿到"页码序号"和"页对象"本身
            for page_index, page in enumerate(pdf):
                text = page.get_text("text").strip()  # 从这一页提取纯文本，去掉首尾空格
                if text:                              # 这页有字才留着，空白页直接跳过
                    documents.append(Document(
                        page_content=text,
                        metadata={
                            "filename": path.name,       # 出处标签：文件名
                            "file_type": "pdf",          # 出处标签：格式
                            "page_number": page_index + 1,  # 页码从 1 开始数（index 从 0 开始，所以 +1）
                        },
                    ))

    # ============ 情况 2：Word 文件，整篇合并 ============
    elif suffix == ".docx":
        doc = WordDocument(path)
        # doc.paragraphs 是 Word 里所有段落的列表；
        # p.text 取段落文字；if p.text.strip() 跳过空段落；
        # "\n".join(...) 用换行把所有段落拼成一个长文本
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if text:
            documents.append(Document(
                page_content=text,
                metadata={"filename": path.name, "file_type": "docx"},  # Word 整篇一张卡，没有页码
            ))

    # ============ 情况 3：Markdown / TXT，纯文本整篇读 ============
    elif suffix in {".md", ".txt"}:
        # utf-8-sig：兼容 Windows 保存的带 BOM 标记的文本文件
        text = path.read_text(encoding="utf-8-sig").strip()
        if text:
            documents.append(Document(
                page_content=text,
                metadata={"filename": path.name, "file_type": suffix.lstrip(".")},
                # suffix.lstrip(".")：把 ".md" 变成 "md"
            ))

    # ============ 情况 4：Excel，一个工作表一张卡 ============
    elif suffix == ".xlsx":
        # read_only=True：大文件读起来更快；data_only=True：读公式的"计算结果"而不是公式本身
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:      # 遍历每个工作表
                rows = []                          # 这个工作表的文本行先放这里
                for row in sheet.iter_rows(values_only=True):  # 逐行取"单元格值"
                    # row 是一行单元格的值元组；
                    # 逐个转成字符串、去空格；None 和空字符串过滤掉
                    values = [
                        str(v).strip()
                        for v in row
                        if v is not None and str(v).strip()
                    ]
                    if values:                     # 这行不是空行
                        rows.append(" | ".join(values))  # 一行表格拼成：报销项目 | 金额
                if rows:                               # 这个工作表有内容
                    documents.append(Document(
                        page_content="\n".join(rows),      # 所有行拼成一段长文本
                        metadata={
                            "filename": path.name,
                            "file_type": "xlsx",
                            "sheet_name": sheet.title,     # 出处标签：工作表名
                        },
                    ))
        finally:
            workbook.close()   # finally：不管成功还是报错，工作簿都一定关闭

    # ============ 情况 5：不认识的后缀，报错 ============
    else:
        raise ValueError(f"不支持的文件格式：{suffix}")

    # ============ 兜底：文件打开了，但一个字都没提取出来，也报错 ============
    if not documents:
        raise ValueError("文件没有提取到可用文本")

    return documents
